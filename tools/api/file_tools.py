"""File access tools for TouriBot chat.

Gives Touri the ability to list and read files from configured knowledge
source folders during conversation.
"""
from __future__ import annotations

import json
from pathlib import Path

FILE_TOOLS = [
    {
        "name": "list_files",
        "description": (
            "List files available in one of Hermann's project folders. "
            "Use this to discover what files exist before reading them. "
            "Available sources: 'Business Wiki', 'Marketing Materials', 'Development Docs'."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "source": {
                    "type": "string",
                    "description": "Source folder label — one of: 'Business Wiki', 'Marketing Materials', 'Development Docs'",
                },
                "search": {
                    "type": "string",
                    "description": "Optional filename pattern to filter (e.g. 'hubspot' or '.xlsx')",
                },
            },
            "required": ["source"],
        },
    },
    {
        "name": "read_file",
        "description": (
            "Read the content of a specific file from a project folder. "
            "Supports text files (.txt, .md, .csv, .json, .html), documents (.pdf, .docx), "
            "and spreadsheets (.xlsx, .csv). For spreadsheets, returns the data as a table. "
            "Use list_files first to find the exact path."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "source": {
                    "type": "string",
                    "description": "Source folder label — one of: 'Business Wiki', 'Marketing Materials', 'Development Docs'",
                },
                "path": {
                    "type": "string",
                    "description": "Relative path within the source folder (from list_files output)",
                },
            },
            "required": ["source", "path"],
        },
    },
]


PROJECT_ROOT = Path(__file__).parent.parent.parent


def _load_sources() -> list[dict]:
    import yaml
    settings_path = PROJECT_ROOT / "args" / "settings.yaml"
    try:
        with open(settings_path) as f:
            cfg = yaml.safe_load(f) or {}
        return cfg.get("knowledge", {}).get("sources", [])
    except Exception:
        return []


def _find_source(label: str) -> dict | None:
    for src in _load_sources():
        if src.get("label", "").lower() == label.lower():
            return src
    return None


def _resolve_root(raw_path: str) -> Path:
    p = Path(raw_path).expanduser()
    if not p.is_absolute():
        p = PROJECT_ROOT / p
    return p.resolve()


def _read_xlsx(filepath: Path) -> str:
    """Read an Excel file and return as markdown table."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(cell) if cell is not None else "" for cell in row])
        if not rows:
            return "(Empty spreadsheet)"
        # Format as markdown table
        # Find actual header row (first row with content)
        header_idx = 0
        for i, row in enumerate(rows):
            if any(cell.strip() for cell in row):
                header_idx = i
                break
        header = rows[header_idx]
        lines = ["| " + " | ".join(header) + " |"]
        lines.append("| " + " | ".join("---" for _ in header) + " |")
        for row in rows[header_idx + 1:]:
            if any(cell.strip() for cell in row):
                lines.append("| " + " | ".join(row) + " |")
        result = "\n".join(lines)
        if len(result) > 20000:
            result = result[:20000] + "\n\n[... truncated at 20,000 characters ...]"
        return result
    except Exception as e:
        return f"Error reading Excel file: {e}"


def _read_csv(filepath: Path) -> str:
    """Read a CSV file and return as markdown table."""
    try:
        import csv
        with open(filepath, newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = list(reader)
        if not rows:
            return "(Empty CSV)"
        header = rows[0]
        lines = ["| " + " | ".join(header) + " |"]
        lines.append("| " + " | ".join("---" for _ in header) + " |")
        for row in rows[1:]:
            if any(cell.strip() for cell in row):
                # Pad row to header length
                padded = row + [""] * (len(header) - len(row))
                lines.append("| " + " | ".join(padded[:len(header)]) + " |")
        result = "\n".join(lines)
        if len(result) > 20000:
            result = result[:20000] + "\n\n[... truncated at 20,000 characters ...]"
        return result
    except Exception as e:
        return f"Error reading CSV: {e}"


def _read_pdf(filepath: Path) -> str:
    """Read a PDF and return extracted text."""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages[:50]:  # Cap at 50 pages
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        result = "\n\n".join(text_parts)
        if len(result) > 20000:
            result = result[:20000] + "\n\n[... truncated at 20,000 characters ...]"
        return result or "(No text extracted from PDF)"
    except Exception as e:
        return f"Error reading PDF: {e}"


def _read_docx(filepath: Path) -> str:
    """Read a DOCX and return extracted text."""
    try:
        import docx
        doc = docx.Document(filepath)
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if len(text) > 20000:
            text = text[:20000] + "\n\n[... truncated at 20,000 characters ...]"
        return text or "(Empty document)"
    except Exception as e:
        return f"Error reading DOCX: {e}"


def handle_file_tool_call(tool_name: str, tool_input: dict) -> str:
    """Execute a file access tool."""
    if tool_name == "list_files":
        source_label = tool_input.get("source", "")
        search = tool_input.get("search", "")

        source = _find_source(source_label)
        if not source:
            available = [s.get("label", "") for s in _load_sources()]
            return f"Source '{source_label}' not found. Available sources: {available}"

        root = _resolve_root(source["path"])
        if not root.exists():
            return f"Source folder does not exist: {root}"

        # List all files recursively
        files = []
        for f in sorted(root.rglob("*")):
            if not f.is_file():
                continue
            if f.name.startswith("."):
                continue
            # Skip .git directories
            if ".git" in f.parts:
                continue
            rel = str(f.relative_to(root))
            if search and search.lower() not in rel.lower():
                continue
            files.append(f"{rel}  ({f.stat().st_size:,} bytes)")

        if not files:
            return f"No files found in '{source_label}'" + (f" matching '{search}'" if search else "")

        result = f"Files in '{source_label}' ({len(files)} files):\n\n"
        result += "\n".join(files[:100])  # Cap at 100 entries
        if len(files) > 100:
            result += f"\n\n... and {len(files) - 100} more files"
        return result

    elif tool_name == "read_file":
        source_label = tool_input.get("source", "")
        rel_path = tool_input.get("path", "")

        source = _find_source(source_label)
        if not source:
            available = [s.get("label", "") for s in _load_sources()]
            return f"Source '{source_label}' not found. Available sources: {available}"

        root = _resolve_root(source["path"])
        if not root.exists():
            return f"Source folder does not exist: {root}"

        # Path traversal protection
        if Path(rel_path).is_absolute():
            return "Error: absolute paths not allowed"

        target = (root / rel_path).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            return "Error: path traversal detected — access denied"

        if not target.exists():
            return f"File not found: {rel_path}"
        if not target.is_file():
            return f"Not a file: {rel_path}"

        # Size check (20MB cap)
        size = target.stat().st_size
        if size > 20 * 1024 * 1024:
            return f"File too large ({size:,} bytes). Maximum: 20MB."

        # Read based on extension
        suffix = target.suffix.lower()

        if suffix == ".xlsx":
            return _read_xlsx(target)
        elif suffix == ".csv":
            return _read_csv(target)
        elif suffix == ".pdf":
            return _read_pdf(target)
        elif suffix in (".docx", ".doc"):
            return _read_docx(target)
        elif suffix in (".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".ico"):
            return f"(Image file: {target.name}, {size:,} bytes — cannot display as text)"
        else:
            # Plain text read
            try:
                content = target.read_text(encoding="utf-8", errors="replace")
                if len(content) > 20000:
                    content = content[:20000] + "\n\n[... truncated at 20,000 characters ...]"
                return content or "(Empty file)"
            except Exception as e:
                return f"Error reading file: {e}"

    return f"Unknown file tool: {tool_name}"
